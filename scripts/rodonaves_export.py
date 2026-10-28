#!/usr/bin/env python3
"""
Baixa o relatório de Frete do Portal do Cliente da Rodonaves e grava na
aba RodonavesRelatorio da Dash_fedex. O Apps Script percebe a mudança e
concilia sozinho: frete (V) e valor da NF (W) das linhas da Rodonaves.

Fluxo (o mesmo que é feito na mão):
  1. rodonaves.com.br/portaldocliente → Entrar → CNPJ + senha
  2. Documentos → Frete (rodonaves.com.br/documentos/frete)
  3. Período (últimos 30 dias, a partir do início da Rodonaves), Buscar
  4. Exportar planilha → pedidos-AAAAMMDD-AAAAMMDD.xlsx

Uma busca só: o "Exportar planilha" baixa TODOS os resultados, não só a
página da tela (conferido na mão com 10 registros por página).

O login é protegido por reCAPTCHA invisível. Pra parecer um navegador
normal, no GitHub o robô roda com a janela visível numa tela virtual
(xvfb) e sem as marcas de automação do Chrome.

Variáveis de ambiente:
  RTE_PORTAL_SENHA   senha do Portal do Cliente
  RTE_PORTAL_CNPJ    (opcional) CNPJ do login — padrão: o da Luuna
  GOOGLE_CREDS_JSON  conteúdo do JSON da service account
  SHEET_ID           (opcional) id da planilha
  RTE_DIAS           (opcional) quantos dias buscar — padrão 30

Rodar local:
  python scripts/rodonaves_export.py --headed --so-baixar
"""

import argparse
import json
import os
import re
import sys
import time
from datetime import date, timedelta
from pathlib import Path

from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout

# ── Configuração ────────────────────────────────────────────────

PORTAL = "https://rodonaves.com.br/portaldocliente"
CNPJ_PADRAO = "42418313000104"
SHEET_ID_PADRAO = "1RMM3Niaf3E9WVCN3NdfU1Cb4Fuz4VSJjaLUPBdrFldI"
ABA_DESTINO = "RodonavesRelatorio"

DIAS_PADRAO = 30            # o que a tela de Frete mostra por padrão
# Primeiro dia com envios pela Rodonaves: antes disso não há o que buscar
INICIO_RODONAVES = date(2026, 10, 27)
TIMEOUT_MS = 60_000
TIMEOUT_DOWNLOAD_MS = 120_000

DIR_SAIDA = Path("saida")
DIR_DEBUG = Path("debug")


def log(msg):
    print(f"[{time.strftime('%H:%M:%S')}] {msg}", flush=True)


def salvar_debug(page, nome):
    DIR_DEBUG.mkdir(exist_ok=True)
    try:
        page.screenshot(path=str(DIR_DEBUG / f"rte_{nome}.png"), full_page=True)
        (DIR_DEBUG / f"rte_{nome}.html").write_text(page.content(), encoding="utf-8")
        log(f"Debug salvo em debug/rte_{nome}.png e .html")
    except Exception as e:
        log(f"Não consegui salvar debug: {e}")


def _so_digitos(s):
    return re.sub(r"\D", "", str(s or ""))


def _primeiro_visivel(page, seletores, timeout=1500):
    """Devolve o primeiro elemento visível entre vários seletores possíveis."""
    for sel in seletores:
        try:
            loc = page.locator(sel)
            for i in range(min(loc.count(), 5)):
                el = loc.nth(i)
                if el.is_visible(timeout=timeout):
                    return el
        except Exception:
            continue
    return None


def _digitar(campo, valor, conferir=None, tentativas=3):
    """Digita num campo com máscara e confere o que entrou."""
    for _ in range(tentativas):
        campo.click()
        campo.press("Control+a")
        campo.press("Delete")
        campo.type(valor, delay=60)
        campo.page.wait_for_timeout(300)
        atual = campo.input_value()
        if conferir is None or conferir(atual):
            return atual
    return campo.input_value()


# ── Etapa 1: login ──────────────────────────────────────────────

def aceitar_cookies(page):
    botao = _primeiro_visivel(page, ["button:has-text('Aceitar')"], timeout=800)
    if botao:
        try:
            botao.click()
            page.wait_for_timeout(500)
        except Exception:
            pass


def fazer_login(page, cnpj, senha):
    log(f"Abrindo {PORTAL}")
    page.goto(PORTAL, wait_until="domcontentloaded", timeout=TIMEOUT_MS)
    page.wait_for_timeout(2500)
    aceitar_cookies(page)

    # Botão "Entrar" do cabeçalho abre o modal "Portal do Cliente"
    entrar = _primeiro_visivel(page, [
        "header button:has-text('Entrar')",
        "header a:has-text('Entrar')",
        "button:has-text('Entrar')",
        "a:has-text('Entrar')",
    ])
    if not entrar:
        salvar_debug(page, "sem_botao_entrar")
        raise RuntimeError("Não achei o botão Entrar no portal — veja debug/")
    entrar.click()
    page.wait_for_timeout(1500)

    campo_senha = _primeiro_visivel(page, ["input[type='password']"], timeout=5000)
    if not campo_senha:
        salvar_debug(page, "sem_modal_login")
        raise RuntimeError("O modal de login não abriu — veja debug/")

    # O CNPJ é o campo de texto do mesmo modal, antes da senha
    campo_cnpj = _primeiro_visivel(page, [
        "[role='dialog'] input:not([type='password']):not([type='hidden']):not([type='checkbox'])",
        "input[placeholder*='CNPJ' i]",
        "input[name*='cnpj' i]",
        "input[id*='cnpj' i]",
    ])
    if not campo_cnpj:
        salvar_debug(page, "sem_campo_cnpj")
        raise RuntimeError("Não achei o campo de CNPJ — veja debug/")

    _digitar(campo_cnpj, cnpj, conferir=lambda v: _so_digitos(v) == cnpj)
    if _so_digitos(campo_cnpj.input_value()) != cnpj:
        salvar_debug(page, "cnpj_incompleto")
        raise RuntimeError("O CNPJ não entrou inteiro no campo — veja debug/")

    _digitar(campo_senha, senha, conferir=lambda v: len(v) == len(senha))
    if len(campo_senha.input_value()) != len(senha):
        salvar_debug(page, "senha_incompleta")
        raise RuntimeError("A senha não entrou inteira no campo — veja debug/")

    # Botão Entrar DO MODAL (o do cabeçalho também se chama Entrar)
    botao = _primeiro_visivel(page, [
        "[role='dialog'] button:has-text('Entrar')",
        "form button[type='submit']",
        "button[type='submit']:has-text('Entrar')",
    ])
    if not botao:
        salvar_debug(page, "sem_botao_modal")
        raise RuntimeError("Não achei o botão Entrar do modal — veja debug/")

    log("Entrando")
    botao.click()

    # Logado = o modal de senha sumiu E apareceu o menu lateral da área
    # logada. Não dá pra usar a palavra "Documentos": a página de antes do
    # login também tem um quadro "Documentos" — foi o que fez o robô achar
    # que já tinha entrado.
    limite = time.time() + 45
    while time.time() < limite:
        page.wait_for_timeout(1000)
        senha_visivel = _primeiro_visivel(page, ["input[type='password']"], timeout=300)
        menu_logado = _primeiro_visivel(page, [
            "text=Minhas coletas",
            "text=Solicitação de coleta",
            "a[href*='/documentos/frete']",
        ], timeout=300)
        if not senha_visivel and menu_logado:
            break
    else:
        salvar_debug(page, "login_falhou")
        texto = ""
        try:
            texto = page.locator("[role='dialog']").inner_text(timeout=1000)[:300]
        except Exception:
            pass
        raise RuntimeError(
            "O login não completou. Pode ser senha errada ou o reCAPTCHA "
            "barrando o robô. Texto do modal: " + (texto or "(vazio)") + " — veja debug/"
        )
    log("Login concluído")


# ── Etapa 2: tela de Frete ──────────────────────────────────────

def abrir_frete(page):
    """Vai pra tela de Frete CLICANDO no menu, como na mão.

    Abrir o endereço /documentos/frete direto dá "página não existe": o
    site só reconhece essa rota quando se navega até ela por dentro.
    """
    log("Abrindo Documentos → Frete pelo menu")

    link = _primeiro_visivel(page, ["a[href*='/documentos/frete']"], timeout=1500)
    if not link:
        # Submenu fechado: abre "Documentos" no menu lateral primeiro
        docs = _primeiro_visivel(page, [
            "nav >> text=Documentos",
            "aside >> text=Documentos",
            "[class*='menu' i] >> text=Documentos",
        ], timeout=1500)
        if docs:
            docs.click()
            page.wait_for_timeout(800)
        link = _primeiro_visivel(page, [
            "a[href*='/documentos/frete']",
            "nav >> text=Frete",
            "aside >> text=Frete",
        ], timeout=3000)

    if not link:
        salvar_debug(page, "sem_menu_frete")
        raise RuntimeError("Não achei Documentos → Frete no menu — veja debug/")

    link.click()
    page.wait_for_timeout(2500)
    if not _primeiro_visivel(page, ["button:has-text('Buscar')"], timeout=10000):
        salvar_debug(page, "sem_tela_frete")
        raise RuntimeError("A tela de Frete não abriu — veja debug/")


def _campos_de_data(page):
    """Os dois campos do Período: os primeiros inputs com dd/mm/aaaa."""
    achados = []
    loc = page.locator("input")
    for i in range(min(loc.count(), 30)):
        el = loc.nth(i)
        try:
            if not el.is_visible(timeout=200):
                continue
            v = el.input_value()
            ph = el.get_attribute("placeholder") or ""
            if re.match(r"^\d{2}/\d{2}/\d{4}$", v) or "dd/mm" in ph.lower():
                achados.append(el)
        except Exception:
            continue
        if len(achados) == 2:
            break
    return achados


def definir_periodo(page, inicio, fim):
    campos = _campos_de_data(page)
    if len(campos) < 2:
        salvar_debug(page, "sem_campos_data")
        raise RuntimeError("Não achei os campos de Período — veja debug/")

    for campo, d in ((campos[0], inicio), (campos[1], fim)):
        alvo = d.strftime("%d/%m/%Y")
        # O campo tem máscara: digitar só os números costuma ser o mais seguro
        atual = _digitar(campo, d.strftime("%d%m%Y"), conferir=lambda v: v == alvo)
        if atual != alvo:
            atual = _digitar(campo, alvo, conferir=lambda v: v == alvo)
        if atual != alvo:
            salvar_debug(page, "data_nao_entrou")
            raise RuntimeError(f"Não consegui preencher a data {alvo} (ficou '{atual}') — veja debug/")
        campo.press("Tab")
        page.wait_for_timeout(200)


def buscar(page):
    page.locator("button:has-text('Buscar')").first.click()
    # Espera a tabela ou o aviso de "nenhum resultado"
    limite = time.time() + 40
    while time.time() < limite:
        page.wait_for_timeout(1000)
        if _primeiro_visivel(page, ["button:has-text('Exportar planilha')"], timeout=300):
            return True
        if _primeiro_visivel(page, ["text=/nenhum (registro|resultado|documento)/i"], timeout=300):
            return False
    return _primeiro_visivel(page, ["button:has-text('Exportar planilha')"], timeout=1000) is not None


def exportar(page, nome):
    DIR_SAIDA.mkdir(exist_ok=True)
    botao = page.locator("button:has-text('Exportar planilha')").first
    with page.expect_download(timeout=TIMEOUT_DOWNLOAD_MS) as dl:
        botao.click()
    arq = DIR_SAIDA / nome
    dl.value.save_as(str(arq))
    return arq


def ler_xlsx(caminho):
    from openpyxl import load_workbook
    ws = load_workbook(caminho, read_only=True, data_only=True).active
    linhas = [list(r) for r in ws.iter_rows(values_only=True)]
    linhas = [l for l in linhas if any(v not in (None, "") for v in l)]
    if not linhas:
        return [], []
    cab = ["" if v is None else str(v).strip() for v in linhas[0]]
    dados = [["" if v is None else v for v in l] for l in linhas[1:]]
    return cab, dados


def baixar_periodo(page, inicio, fim):
    rotulo = f"{inicio:%d/%m/%Y} a {fim:%d/%m/%Y}"
    definir_periodo(page, inicio, fim)
    if not buscar(page):
        log(f"{rotulo}: nenhum CT-e")
        return None, []
    arq = exportar(page, f"rte_{inicio:%Y%m%d}_{fim:%Y%m%d}.xlsx")
    cab, dados = ler_xlsx(arq)
    log(f"{rotulo}: {len(dados)} CT-e(s)")
    return cab, dados


# ── Etapa 3: planilha ───────────────────────────────────────────

def _indice(cab, *nomes):
    norm = [re.sub(r"\s+", " ", c).strip().lower() for c in cab]
    for n in nomes:
        n = n.lower()
        if n in norm:
            return norm.index(n)
    return None


def escrever_no_sheets(cab, dados, sheet_id, creds_json):
    import gspread
    from google.oauth2.service_account import Credentials

    if not dados:
        raise RuntimeError("Nenhuma linha baixada — abortando pra não mexer na aba")

    # Chave de mescla: o CT-e (uma NF pode ter mais de um, e queremos todos)
    i_chave = _indice(cab, "chave de acesso ct-e")
    if i_chave is None:
        i_chave = _indice(cab, "ct-e")
    if i_chave is None:
        raise RuntimeError(f"Não achei a coluna do CT-e no relatório. Cabeçalho: {cab}")

    def chave(l):
        return _so_digitos(l[i_chave]) if i_chave < len(l) else ""

    log("Conectando no Google Sheets")
    creds = Credentials.from_service_account_info(
        json.loads(creds_json), scopes=["https://www.googleapis.com/auth/spreadsheets"])
    planilha = gspread.authorize(creds).open_by_key(sheet_id)
    try:
        aba = planilha.worksheet(ABA_DESTINO)
    except gspread.WorksheetNotFound:
        log(f'Aba "{ABA_DESTINO}" não existe — criando')
        aba = planilha.add_worksheet(title=ABA_DESTINO, rows=1000, cols=len(cab) + 2)

    # O portal só mostra 30 dias: mescla com o que já existe, pra não
    # perder o histórico mais antigo. Lê sem formatação, pra os valores
    # continuarem número.
    try:
        existentes = aba.get_values(value_render_option="UNFORMATTED_VALUE")
    except Exception:
        existentes = []

    por_cte = {}
    if existentes and len(existentes) >= 2 and [str(c).strip() for c in existentes[0]] == cab:
        for l in existentes[1:]:
            k = chave(l)
            if k:
                por_cte[k] = l
    elif existentes and len(existentes) >= 2:
        log("Cabeçalho da aba diferente do relatório — regravando do zero")

    antes, atualizadas = len(por_cte), 0
    for l in dados:
        k = chave(l)
        if not k:
            continue
        if k in por_cte:
            atualizadas += 1
        por_cte[k] = l

    finais = [[("" if v is None else v) for v in (l + [""] * len(cab))[:len(cab)]]
              for l in por_cte.values()]
    log(f"Mesclado — {antes} já existiam, {atualizadas} atualizado(s), "
        f"{len(finais) - antes} novo(s), {len(finais)} no total")

    aba.clear()
    # RAW: chaves de 44 dígitos e números de pedido ficam como texto, e
    # os valores (que já vêm como número do Excel) continuam número
    aba.update(values=[cab] + finais, range_name="A1", value_input_option="RAW")
    log(f'Aba "{ABA_DESTINO}" atualizada com {len(finais)} linha(s)')
    return len(finais)


# ── Principal ───────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--headed", action="store_true", help="mostra o navegador")
    ap.add_argument("--so-baixar", action="store_true", help="não escreve na planilha")
    ap.add_argument("--dias", type=int, default=None)
    args = ap.parse_args()

    senha = os.environ.get("RTE_PORTAL_SENHA")
    cnpj = _so_digitos(os.environ.get("RTE_PORTAL_CNPJ") or CNPJ_PADRAO)
    sheet_id = os.environ.get("SHEET_ID", SHEET_ID_PADRAO)
    creds_json = os.environ.get("GOOGLE_CREDS_JSON")
    dias = args.dias or int(os.environ.get("RTE_DIAS", DIAS_PADRAO))

    if not senha:
        sys.exit("Falta a variável RTE_PORTAL_SENHA")
    if not args.so_baixar and not creds_json:
        sys.exit("Falta a variável GOOGLE_CREDS_JSON")
    log(f"CNPJ {cnpj[:2]}.***.***/{cnpj[8:12]}-** | senha: {len(senha)} caractere(s) | {dias} dia(s)")

    hoje = date.today()
    inicio_total = max(hoje - timedelta(days=dias - 1), INICIO_RODONAVES)

    with sync_playwright() as p:
        navegador = p.chromium.launch(
            headless=not args.headed,
            args=["--disable-blink-features=AutomationControlled"],
        )
        contexto = navegador.new_context(
            viewport={"width": 1600, "height": 950},
            accept_downloads=True,
            locale="pt-BR",
            timezone_id="America/Sao_Paulo",
            user_agent=("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                        "(KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36"),
        )
        # Tira a marca mais óbvia de automação, que o reCAPTCHA olha
        contexto.add_init_script("Object.defineProperty(navigator, 'webdriver', {get: () => undefined})")
        page = contexto.new_page()

        try:
            fazer_login(page, cnpj, senha)
            abrir_frete(page)

            log(f"Buscando de {inicio_total:%d/%m/%Y} a {hoje:%d/%m/%Y}")
            cab, todos = baixar_periodo(page, inicio_total, hoje)
            if args.so_baixar:
                log("Modo --so-baixar: não escrevi na planilha")
            elif not cab:
                log("Nenhum CT-e no período — nada a gravar")
            else:
                escrever_no_sheets(cab, todos, sheet_id, creds_json)
                log("Concluído")

        except PWTimeout as e:
            salvar_debug(page, "timeout")
            sys.exit(f"Timeout: {e}")
        except Exception as e:
            salvar_debug(page, "erro")
            sys.exit(f"Erro: {e}")
        finally:
            contexto.close()
            navegador.close()


if __name__ == "__main__":
    main()
